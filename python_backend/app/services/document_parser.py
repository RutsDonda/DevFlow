"""
app/services/document_parser.py – Extract plain text from uploaded files.

Supported: PDF (.pdf), Word (.docx) and text-like files (.txt, .md, .csv, .json,
source code …). Legacy binary formats (.doc) are rejected with a helpful message.
"""

import io

from fastapi import HTTPException, status

TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "rst", "csv", "tsv", "json", "yaml", "yml", "xml", "html", "htm",
    "log", "ini", "toml", "py", "js", "jsx", "ts", "tsx", "java", "c", "h", "cpp", "hpp", "cs",
    "go", "rs", "rb", "php", "sql", "sh", "kt", "swift",
}
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _bad(msg: str) -> HTTPException:
    return HTTPException(status.HTTP_400_BAD_REQUEST, msg)


def extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Return (text, content_type). Raises HTTP 400 for unsupported or empty files. Blocking."""
    ext = extension(filename)

    if ext == "pdf":
        text, mime = _pdf(data), PDF_MIME
    elif ext == "docx":
        text, mime = _docx(data), DOCX_MIME
    elif ext == "doc":
        raise _bad("Legacy .doc files are not supported – save the file as .docx and upload it again")
    elif ext in TEXT_EXTENSIONS or (not ext and b"\x00" not in data[:4096]):
        text, mime = data.decode("utf-8", errors="replace"), "text/plain"
    else:
        raise _bad(f"Unsupported file type '.{ext}'. Upload a PDF, Word (.docx) or text file.")

    text = text.replace("\x00", "").strip()
    if not text:
        raise _bad(f"'{filename}' has no extractable text (scanned or image-only documents need OCR first)")
    return text, mime


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:
        raise _bad(f"Could not read PDF: {exc}")
    return "\n\n".join(f"[Page {i}]\n{p}" for i, p in enumerate(pages, 1) if p)


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise _bad(f"Could not read Word document: {exc}")

    parts = []
    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        # Keep headings recognisable so chunks carry structure
        style = (para.style.name or "").lower() if para.style is not None else ""
        parts.append(f"## {para.text}" if style.startswith("heading") else para.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)
